"""JSON và bảng mẫu openpyxl theo GUIDE; không đưa trọng số vào JSON."""
from __future__ import annotations
from copy import copy
import json
import math
import os
from pathlib import Path


def save_result(result, results_dir="../results"):
    path = Path(results_dir) / f"{result['cfg']['exp_id']}.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {k: v for k, v in result.items() if k != "best_state"}
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    os.replace(tmp, path)
    return str(path)


def load_results(results_dir="../results"):
    results = []
    for path in sorted(Path(results_dir).glob("*.json")):
        item = json.loads(path.read_text(encoding="utf-8"))
        # Part 1 health evidence có schema riêng, không là thí nghiệm Part 2/3.
        if all(k in item for k in ("cfg", "history", "summary")):
            results.append(item)
    return sorted(results, key=lambda r: r["cfg"]["exp_id"])


def to_row(result, eval_scores=None, notes=""):
    c, s = result["cfg"], result["summary"]
    if eval_scores is not None and c["group"] != "baseline" and not result.get("selected_final"):
        raise ValueError("Chỉ điền eval cho baseline hoặc cấu hình đã chọn cuối")
    row = {**c, **s}
    row["hidden"] = "-".join(map(str, c["hidden"]))
    row["clip_norm"] = c["clip_norm"] if c["clip_norm"] is not None else "none"
    row["figure_file"] = f"figures/{c['exp_id']}.png"
    row["notes"] = "; ".join(filter(None, [notes, c.get("notes"), s.get("reason"),
        f"status={s.get('status')}", f"betas={c.get('betas')}; eps={c.get('eps')}",
        f"train monitor={result.get('initial', {}).get('train_monitor_size', 'n.a.')} fixed samples; eval-mode FP32",
        f"clip fraction={s.get('clip_fraction')}; skipped AMP updates={s.get('skipped_updates', 0)}",
        f"activation std={result.get('initial', {}).get('activation_std')}"]))
    row["eval_acc"] = eval_scores["accuracy"] if eval_scores else None
    row["eval_macro_f1"] = eval_scores["macro_f1"] if eval_scores else None
    return row


def seed_statistics(results):
    import numpy as np
    base = [r for r in results if r["cfg"]["group"] == "baseline" and r["summary"].get("status") == "complete"]
    answer = {"n": len(base), "exp_ids": [r["cfg"]["exp_id"] for r in base]}
    for key in ("val_acc", "val_macro_f1", "best_val_loss"):
        values = [r["summary"][key] for r in base]
        answer[key] = dict(mean=float(np.mean(values)) if values else None,
            std=float(np.std(values, ddof=1)) if len(values) >= 2 else None)
    answer["noise_2sigma"] = 2 * answer["val_macro_f1"]["std"] if len(base) >= 2 else None
    return answer


def write_xlsx(rows, template_path, out_path):
    """Giữ 4 sheet, tiêu đề và công thức; mở Excel/LibreOffice để tính lại.

    Mẫu Seeds có 5 slot, đủ baseline mặc định 3 seed. Phạm vi công thức
    được mở rộng nếu số thí nghiệm vượt 60. Không ghi đè cột công thức bằng số.
    """
    import openpyxl
    from openpyxl.formula.translate import Translator
    from openpyxl.workbook.properties import CalcProperties
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]
    headers = [cell.value for cell in ws[1]]
    formula_names = {"step0_gap_vs_lnC", "gap_val_minus_train", "delta_val_f1_vs_base", "beyond_noise"}
    end = max(61, len(rows) + 1)
    for i in range(2, end + 1):
        row = rows[i - 2] if i - 2 < len(rows) else {}
        for col, name in enumerate(headers, 1):
            cell = ws.cell(i, col)
            if i > 61:
                cell._style = copy(ws.cell(2, col)._style)
                cell.alignment = copy(ws.cell(2, col).alignment)
            if name in formula_names:
                if not cell.value:
                    cell.value = Translator(ws.cell(2, col).value, origin=ws.cell(2, col).coordinate).translate_formula(cell.coordinate)
            else:
                value = row.get(name)
                if isinstance(value, float) and not math.isfinite(value):
                    value = None
                cell.value = value
    seeds = wb["Seeds"]
    base = [r for r in rows if r["group"] == "baseline" and r.get("status") == "complete"]
    if len(base) > 5:
        raise ValueError("Mẫu có 5 slot Seeds; chọn tối đa 5 baseline seed")
    for i in range(2, 7):
        seeds.cell(i, 1).value = base[i - 2]["exp_id"] if i - 2 < len(base) else None
    for sheet in (seeds, wb["Summary"]):
        for line in sheet:
            for cell in line:
                if cell.data_type == "f" and end > 61:
                    cell.value = cell.value.replace("$61", f"${end}")
    summary = wb["Summary"]
    for i in range(2, 11):
        group = summary.cell(i, 1).value
        members = [(j + 2, r) for j, r in enumerate(rows) if r["group"] == group and isinstance(r.get("val_macro_f1"), (int, float))]
        # Thay MAXIFS/MINIFS của mẫu bằng MAX/MIN trên các ô cùng nhóm:
        # tương thích cả Excel cũ, không thay ý nghĩa và vẫn là công thức.
        for target_col, source_col, operation in ((4, "V", "MAX"), (5, "V", "MIN"), (6, "U", "MAX")):
            refs = ",".join(f"Experiments!{source_col}{j}" for j, _ in members)
            summary.cell(i, target_col).value = f"={operation}({refs})" if refs else '=IF(1=1,"",0)'
        if members:
            best = max((r for _, r in members), key=lambda r: r["val_macro_f1"])
            summary.cell(i, 8).value = f"Val F1 cao nhất: {best['exp_id']} = {best['val_macro_f1']:.4f}; xem REPORT.md để đối chiếu nhiễu/cơ chế."
        else:
            summary.cell(i, 8).value = "Chưa có kết quả hoàn tất; xem notes nếu phần cứng không hỗ trợ."
    ws.auto_filter.ref = f"A1:AG{len(rows) + 1}"
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    path = Path(out_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
